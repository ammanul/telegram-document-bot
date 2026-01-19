"""
This module provides a function to convert a list of images into a PPTX document.
"""
import io
from typing import List

from pptx import Presentation
from pptx.util import Inches
from PIL import Image

from worker.app.config import (
    EXPORT_PAGE_HEIGHT,
    EXPORT_PAGE_WIDTH,
    EXPORT_PAGE_PADDING_TOP,
    EXPORT_PAGE_PADDING_BOTTOM,
    EXPORT_PAGE_PADDING_LEFT,
    EXPORT_PAGE_PADDING_RIGHT,
    ID_SPACING_Y,
)


def images_to_ppt_stream(images: List[Image.Image]) -> io.BytesIO:
    """
    Converts a list of PIL Images into a PPTX document as a byte stream.

    Args:
        images: A list of PIL Image objects.

    Returns:
        An in-memory byte stream (io.BytesIO) containing the PPTX document.

    Raises:
        ValueError: If any of the provided items is not a valid image or fails during processing.
    """
    prs = Presentation()
    prs.slide_width = Inches(EXPORT_PAGE_WIDTH)
    prs.slide_height = Inches(EXPORT_PAGE_HEIGHT)
    blank_slide_layout = prs.slide_layouts[6]

    left_padding = Inches(EXPORT_PAGE_PADDING_LEFT)
    right_padding = Inches(EXPORT_PAGE_PADDING_RIGHT)
    top_padding = Inches(EXPORT_PAGE_PADDING_TOP)
    bottom_padding = Inches(EXPORT_PAGE_PADDING_BOTTOM)
    gutter = Inches(ID_SPACING_Y)

    available_width = prs.slide_width - left_padding - right_padding
    available_height = prs.slide_height - top_padding - bottom_padding

    current_slide = prs.slides.add_slide(blank_slide_layout)
    height_used_on_page = top_padding

    for i, image in enumerate(images):
        if not isinstance(image, Image.Image):
            raise ValueError(f"Invalid image provided at index {i}")

        try:
            if image.mode == "RGBA":
                background = Image.new("RGB", image.size, (255, 255, 255))
                background.paste(image, (0, 0), image)
                processed_image = background
            else:
                processed_image = image.convert("RGB")

            img_buffer = io.BytesIO()
            processed_image.save(img_buffer, format="JPEG", quality=95)
            img_buffer.seek(0)

            img_width, img_height = image.size
            aspect_ratio = img_height / img_width if img_width > 0 else 1
            new_width = available_width
            new_height = new_width * aspect_ratio

            current_gutter = gutter if i > 0 else 0

            if height_used_on_page + new_height + current_gutter > prs.slide_height - bottom_padding:
                current_slide = prs.slides.add_slide(blank_slide_layout)
                height_used_on_page = top_padding
                current_gutter = 0

            if current_gutter > 0:
                height_used_on_page += current_gutter

            current_slide.shapes.add_picture(
                img_buffer, left_padding, height_used_on_page, width=new_width, height=new_height
            )
            height_used_on_page += new_height

        except Exception as e:
            raise ValueError(f"Failed to process image {i} for PPTX: {e}")

    # Save the presentation to a byte stream
    pptx_stream = io.BytesIO()
    prs.save(pptx_stream)
    pptx_stream.seek(0)
    return pptx_stream
